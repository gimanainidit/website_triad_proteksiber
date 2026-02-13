#!/usr/bin/env python3
"""
TRIAD AI-based CMS Backend Testing Suite
Tests all backend APIs and Supabase integration
"""

import requests
import json
import os
import tempfile
import uuid
from pathlib import Path
import time

# Get backend URL from frontend .env file
def get_backend_url():
    frontend_env_path = Path("/app/frontend/.env")
    if frontend_env_path.exists():
        with open(frontend_env_path, 'r') as f:
            for line in f:
                if line.startswith('REACT_APP_BACKEND_URL='):
                    return line.split('=', 1)[1].strip()
    return "http://localhost:8001"

BASE_URL = get_backend_url()
API_BASE = f"{BASE_URL}/api"

class Colors:
    GREEN = '\033[92m'
    RED = '\033[91m'
    YELLOW = '\033[93m'
    BLUE = '\033[94m'
    ENDC = '\033[0m'
    BOLD = '\033[1m'

def print_test_header(test_name):
    print(f"\n{Colors.BLUE}{Colors.BOLD}{'='*60}{Colors.ENDC}")
    print(f"{Colors.BLUE}{Colors.BOLD}Testing: {test_name}{Colors.ENDC}")
    print(f"{Colors.BLUE}{Colors.BOLD}{'='*60}{Colors.ENDC}")

def print_success(message):
    print(f"{Colors.GREEN}✅ {message}{Colors.ENDC}")

def print_error(message):
    print(f"{Colors.RED}❌ {message}{Colors.ENDC}")

def print_warning(message):
    print(f"{Colors.YELLOW}⚠️  {message}{Colors.ENDC}")

def print_info(message):
    print(f"{Colors.BLUE}ℹ️  {message}{Colors.ENDC}")

def create_test_pdf():
    """Create a simple test PDF file"""
    try:
        from reportlab.pdfgen import canvas
        from reportlab.lib.pagesizes import letter
        
        # Create temporary PDF file
        temp_file = tempfile.NamedTemporaryFile(delete=False, suffix='.pdf')
        
        # Create PDF content
        c = canvas.Canvas(temp_file.name, pagesize=letter)
        c.drawString(100, 750, "Test PDF Document for TRIAD CMS")
        c.drawString(100, 730, "This is a test file for upload functionality")
        c.drawString(100, 710, f"Created at: {time.strftime('%Y-%m-%d %H:%M:%S')}")
        c.save()
        
        return temp_file.name
    except ImportError:
        # Fallback: create a simple text file with PDF extension
        temp_file = tempfile.NamedTemporaryFile(delete=False, suffix='.pdf', mode='w')
        temp_file.write("%PDF-1.4\n")
        temp_file.write("1 0 obj\n")
        temp_file.write("<<\n")
        temp_file.write("/Type /Catalog\n")
        temp_file.write("/Pages 2 0 R\n")
        temp_file.write(">>\n")
        temp_file.write("endobj\n")
        temp_file.write("Test PDF content for TRIAD CMS upload testing")
        temp_file.close()
        return temp_file.name

def create_test_pptx():
    """Create a simple test PPTX file"""
    try:
        from pptx import Presentation
        from pptx.util import Inches
        
        # Create temporary PPTX file
        temp_file = tempfile.NamedTemporaryFile(delete=False, suffix='.pptx')
        temp_file.close()
        
        # Create PowerPoint presentation
        prs = Presentation()
        
        # Add first slide
        slide_layout = prs.slide_layouts[0]  # Title slide
        slide = prs.slides.add_slide(slide_layout)
        title = slide.shapes.title
        subtitle = slide.placeholders[1]
        
        title.text = "Test PPTX Document"
        subtitle.text = "This is a test PowerPoint file for TRIAD CMS conversion testing"
        
        # Add second slide with content
        slide_layout = prs.slide_layouts[1]  # Title and content
        slide = prs.slides.add_slide(slide_layout)
        title = slide.shapes.title
        content = slide.placeholders[1]
        
        title.text = "Test Content Slide"
        content.text = "This slide contains test content that should be converted to PDF format."
        
        # Save the presentation
        prs.save(temp_file.name)
        
        return temp_file.name
    except ImportError:
        print_warning("python-pptx not available, creating mock PPTX file")
        # Create a mock PPTX file (actually a ZIP with minimal structure)
        import zipfile
        temp_file = tempfile.NamedTemporaryFile(delete=False, suffix='.pptx')
        temp_file.close()
        
        with zipfile.ZipFile(temp_file.name, 'w') as zf:
            zf.writestr('[Content_Types].xml', '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>')
            zf.writestr('_rels/.rels', '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>')
        
        return temp_file.name

def create_oversized_file():
    """Create a file larger than 15MB for testing size limits"""
    temp_file = tempfile.NamedTemporaryFile(delete=False, suffix='.pdf', mode='wb')
    
    # Write 16MB of data
    chunk_size = 1024 * 1024  # 1MB chunks
    for i in range(16):
        temp_file.write(b'0' * chunk_size)
    
    temp_file.close()
    return temp_file.name

def test_health_check():
    """Test API health check endpoint"""
    print_test_header("API Health Check")
    
    try:
        response = requests.get(f"{API_BASE}/health", timeout=10)
        
        if response.status_code == 200:
            data = response.json()
            print_success(f"Health check passed: {data}")
            return True
        else:
            print_error(f"Health check failed with status {response.status_code}: {response.text}")
            return False
            
    except requests.exceptions.RequestException as e:
        print_error(f"Health check request failed: {str(e)}")
        return False

def test_get_resources_empty():
    """Test GET /api/resources - should return empty array initially"""
    print_test_header("GET Resources (Empty State)")
    
    try:
        response = requests.get(f"{API_BASE}/resources", timeout=10)
        
        if response.status_code == 200:
            data = response.json()
            if isinstance(data, list):
                print_success(f"Resources endpoint working, returned {len(data)} resources")
                return True
            else:
                print_error(f"Expected list, got: {type(data)}")
                return False
        else:
            print_error(f"GET resources failed with status {response.status_code}: {response.text}")
            return False
            
    except requests.exceptions.RequestException as e:
        print_error(f"GET resources request failed: {str(e)}")
        return False

def test_get_categories_empty():
    """Test GET /api/categories - should return empty array initially"""
    print_test_header("GET Categories (Empty State)")
    
    try:
        response = requests.get(f"{API_BASE}/categories", timeout=10)
        
        if response.status_code == 200:
            data = response.json()
            if "categories" in data and isinstance(data["categories"], list):
                print_success(f"Categories endpoint working, returned {len(data['categories'])} categories")
                return True
            else:
                print_error(f"Expected categories list, got: {data}")
                return False
        else:
            print_error(f"GET categories failed with status {response.status_code}: {response.text}")
            return False
            
    except requests.exceptions.RequestException as e:
        print_error(f"GET categories request failed: {str(e)}")
        return False

def test_file_upload():
    """Test POST /api/admin/resources/upload - file upload functionality"""
    print_test_header("File Upload Test")
    
    # Create test PDF file
    test_file_path = create_test_pdf()
    
    try:
        # Prepare form data
        with open(test_file_path, 'rb') as f:
            files = {
                'file': ('test_document.pdf', f, 'application/pdf')
            }
            data = {
                'title': 'Test Resource Document',
                'description': 'This is a test document for TRIAD CMS upload functionality',
                'category': 'Panduan',
                'tags': 'test, upload, pdf, cms'
            }
            
            response = requests.post(
                f"{API_BASE}/admin/resources/upload",
                files=files,
                data=data,
                timeout=30
            )
        
        if response.status_code == 200:
            data = response.json()
            print_success(f"File upload successful!")
            print_info(f"Response: {json.dumps(data, indent=2)}")
            
            # Verify required fields in response
            required_fields = ['id', 'title', 'message', 'file_url']
            for field in required_fields:
                if field not in data:
                    print_error(f"Missing required field in response: {field}")
                    return False, None
            
            return True, data['id']
        else:
            print_error(f"File upload failed with status {response.status_code}: {response.text}")
            return False, None
            
    except requests.exceptions.RequestException as e:
        print_error(f"File upload request failed: {str(e)}")
        return False, None
    finally:
        # Clean up test file
        try:
            os.unlink(test_file_path)
        except:
            pass

def test_get_single_resource(resource_id):
    """Test GET /api/resources/{id} - get single resource"""
    print_test_header(f"GET Single Resource (ID: {resource_id})")
    
    try:
        response = requests.get(f"{API_BASE}/resources/{resource_id}", timeout=10)
        
        if response.status_code == 200:
            data = response.json()
            print_success("Single resource retrieval successful!")
            print_info(f"Resource data: {json.dumps(data, indent=2)}")
            
            # Verify required fields
            required_fields = ['id', 'title', 'description', 'category', 'file_url']
            for field in required_fields:
                if field not in data:
                    print_error(f"Missing required field: {field}")
                    return False
            
            return True
        elif response.status_code == 404:
            print_error("Resource not found (404)")
            return False
        else:
            print_error(f"GET single resource failed with status {response.status_code}: {response.text}")
            return False
            
    except requests.exceptions.RequestException as e:
        print_error(f"GET single resource request failed: {str(e)}")
        return False

def test_get_resources_after_upload():
    """Test GET /api/resources after upload - should contain uploaded resource"""
    print_test_header("GET Resources (After Upload)")
    
    try:
        response = requests.get(f"{API_BASE}/resources", timeout=10)
        
        if response.status_code == 200:
            data = response.json()
            if isinstance(data, list) and len(data) > 0:
                print_success(f"Resources found after upload: {len(data)} resources")
                print_info(f"First resource: {json.dumps(data[0], indent=2)}")
                return True
            else:
                print_error("No resources found after upload")
                return False
        else:
            print_error(f"GET resources failed with status {response.status_code}: {response.text}")
            return False
            
    except requests.exceptions.RequestException as e:
        print_error(f"GET resources request failed: {str(e)}")
        return False

def test_file_url_accessibility(file_url):
    """Test if uploaded file URL is accessible"""
    print_test_header("File URL Accessibility Test")
    
    try:
        response = requests.head(file_url, timeout=10)
        
        if response.status_code == 200:
            print_success(f"File URL is accessible: {file_url}")
            return True
        else:
            print_error(f"File URL not accessible, status {response.status_code}: {file_url}")
            return False
            
    except requests.exceptions.RequestException as e:
        print_error(f"File URL accessibility test failed: {str(e)}")
        return False

def test_delete_resource(resource_id):
    """Test DELETE /api/admin/resources/{id} - delete resource"""
    print_test_header(f"DELETE Resource (ID: {resource_id})")
    
    try:
        response = requests.delete(f"{API_BASE}/admin/resources/{resource_id}", timeout=10)
        
        if response.status_code == 200:
            data = response.json()
            print_success(f"Resource deleted successfully: {data}")
            return True
        elif response.status_code == 404:
            print_error("Resource not found for deletion (404)")
            return False
        else:
            print_error(f"DELETE resource failed with status {response.status_code}: {response.text}")
            return False
            
    except requests.exceptions.RequestException as e:
        print_error(f"DELETE resource request failed: {str(e)}")
        return False

def test_pptx_to_pdf_conversion():
    """Test PPTX to PDF conversion functionality"""
    print_test_header("PPTX to PDF Conversion Test")
    
    # Create test PPTX file
    test_file_path = create_test_pptx()
    
    try:
        # Prepare form data
        with open(test_file_path, 'rb') as f:
            files = {
                'file': ('test_presentation.pptx', f, 'application/vnd.openxmlformats-officedocument.presentationml.presentation')
            }
            data = {
                'title': 'Test PPTX Conversion',
                'description': 'This PPTX file should be converted to PDF',
                'category': 'Presentation',
                'tags': 'test, pptx, conversion, pdf'
            }
            
            response = requests.post(
                f"{API_BASE}/admin/resources/upload",
                files=files,
                data=data,
                timeout=30
            )
        
        if response.status_code == 200:
            data = response.json()
            print_success(f"PPTX upload and conversion successful!")
            print_info(f"Response: {json.dumps(data, indent=2)}")
            
            # Verify conversion indicators
            if 'converted_from' in data and data['converted_from'] == 'pptx':
                print_success("✅ Conversion metadata correctly set")
            else:
                print_error("❌ Missing or incorrect conversion metadata")
                return False, None
            
            if 'PPTX converted to PDF' in data.get('message', ''):
                print_success("✅ Conversion message correctly included")
            else:
                print_error("❌ Missing conversion message")
                return False, None
            
            return True, data['id']
        else:
            print_error(f"PPTX upload failed with status {response.status_code}: {response.text}")
            return False, None
            
    except requests.exceptions.RequestException as e:
        print_error(f"PPTX upload request failed: {str(e)}")
        return False, None
    finally:
        # Clean up test file
        try:
            os.unlink(test_file_path)
        except:
            pass

def test_file_size_limit():
    """Test file size limit (15MB max)"""
    print_test_header("File Size Limit Test (15MB)")
    
    # Create oversized file
    test_file_path = create_oversized_file()
    
    try:
        with open(test_file_path, 'rb') as f:
            files = {
                'file': ('oversized_file.pdf', f, 'application/pdf')
            }
            data = {
                'title': 'Oversized File Test',
                'description': 'This file should be rejected due to size',
                'category': 'Test',
                'tags': 'oversized, test'
            }
            
            response = requests.post(
                f"{API_BASE}/admin/resources/upload",
                files=files,
                data=data,
                timeout=30
            )
        
        if response.status_code == 400:
            response_data = response.json()
            if 'exceeds' in response_data.get('detail', '').lower():
                print_success("Oversized file correctly rejected with size limit message")
                return True
            else:
                print_error(f"File rejected but wrong error message: {response_data}")
                return False
        else:
            print_error(f"Expected 400 for oversized file, got {response.status_code}: {response.text}")
            return False
            
    except requests.exceptions.RequestException as e:
        print_error(f"File size limit test failed: {str(e)}")
        return False
    finally:
        try:
            os.unlink(test_file_path)
        except:
            pass

def test_invalid_file_types():
    """Test various invalid file types"""
    print_test_header("Invalid File Types Test")
    
    invalid_files = [
        ('test.txt', 'text/plain', 'This is a text file'),
        ('test.doc', 'application/msword', 'This is a DOC file'),
        ('test.jpg', 'image/jpeg', b'\xff\xd8\xff\xe0\x00\x10JFIF'),
        ('test.xlsx', 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', 'Excel content')
    ]
    
    all_passed = True
    
    for filename, content_type, content in invalid_files:
        print_info(f"Testing {filename}...")
        
        # Create temporary file
        if isinstance(content, bytes):
            temp_file = tempfile.NamedTemporaryFile(delete=False, mode='wb')
            temp_file.write(content)
        else:
            temp_file = tempfile.NamedTemporaryFile(delete=False, mode='w')
            temp_file.write(content)
        temp_file.close()
        
        try:
            with open(temp_file.name, 'rb') as f:
                files = {
                    'file': (filename, f, content_type)
                }
                data = {
                    'title': f'Invalid File Test - {filename}',
                    'description': 'This should be rejected',
                    'category': 'Test',
                    'tags': 'invalid'
                }
                
                response = requests.post(
                    f"{API_BASE}/admin/resources/upload",
                    files=files,
                    data=data,
                    timeout=10
                )
            
            if response.status_code == 400:
                response_data = response.json()
                if 'not allowed' in response_data.get('detail', '').lower():
                    print_success(f"✅ {filename} correctly rejected")
                else:
                    print_error(f"❌ {filename} rejected but wrong error message: {response_data}")
                    all_passed = False
            else:
                print_error(f"❌ {filename} should be rejected, got {response.status_code}: {response.text}")
                all_passed = False
                
        except requests.exceptions.RequestException as e:
            print_error(f"Request failed for {filename}: {str(e)}")
            all_passed = False
        finally:
            try:
                os.unlink(temp_file.name)
            except:
                pass
    
    return all_passed

def test_database_verification(resource_id):
    """Test database verification for converted files"""
    print_test_header(f"Database Verification (ID: {resource_id})")
    
    try:
        response = requests.get(f"{API_BASE}/resources/{resource_id}", timeout=10)
        
        if response.status_code == 200:
            data = response.json()
            
            # Check file_type is 'pdf'
            if data.get('file_type') == 'pdf':
                print_success("✅ File type correctly stored as 'pdf'")
            else:
                print_error(f"❌ Expected file_type 'pdf', got '{data.get('file_type')}'")
                return False
            
            # Check description includes conversion note
            description = data.get('description', '')
            if 'converted from PPTX' in description:
                print_success("✅ Description includes conversion note")
            else:
                print_error(f"❌ Description missing conversion note: {description}")
                return False
            
            # Check file URL points to PDF
            file_url = data.get('file_url', '')
            if file_url.endswith('.pdf'):
                print_success("✅ File URL points to PDF file")
            else:
                print_error(f"❌ File URL doesn't end with .pdf: {file_url}")
                return False
            
            return True
        else:
            print_error(f"Failed to get resource for verification: {response.status_code}")
            return False
            
    except requests.exceptions.RequestException as e:
        print_error(f"Database verification request failed: {str(e)}")
        return False

def test_pdf_content_type(file_url):
    """Test that stored files have correct PDF content-type"""
    print_test_header("PDF Content-Type Verification")
    
    try:
        response = requests.head(file_url, timeout=10)
        
        if response.status_code == 200:
            content_type = response.headers.get('content-type', '').lower()
            if 'application/pdf' in content_type:
                print_success(f"✅ Correct content-type: {content_type}")
                return True
            else:
                print_error(f"❌ Expected PDF content-type, got: {content_type}")
                return False
        else:
            print_error(f"Failed to check content-type, status: {response.status_code}")
            return False
            
    except requests.exceptions.RequestException as e:
        print_error(f"Content-type check failed: {str(e)}")
        return False

def test_missing_required_fields():
    """Test upload with missing required fields"""
    print_test_header("Missing Required Fields Test")
    
    test_file_path = create_test_pdf()
    
    try:
        with open(test_file_path, 'rb') as f:
            files = {
                'file': ('test_document.pdf', f, 'application/pdf')
            }
            # Missing title and description
            data = {
                'category': 'Test'
            }
            
            response = requests.post(
                f"{API_BASE}/admin/resources/upload",
                files=files,
                data=data,
                timeout=10
            )
        
        if response.status_code == 422:  # FastAPI validation error
            print_success("Missing required fields correctly rejected (422)")
            return True
        else:
            print_error(f"Expected 422 for missing fields, got {response.status_code}: {response.text}")
            return False
            
    except requests.exceptions.RequestException as e:
        print_error(f"Missing required fields test failed: {str(e)}")
        return False
    finally:
        try:
            os.unlink(test_file_path)
        except:
            pass

# AI Generation Testing Functions

def test_ai_generation_types():
    """Test GET /api/ai/generation-types endpoint"""
    print_test_header("AI Generation Types Endpoint")
    
    try:
        response = requests.get(f"{API_BASE}/ai/generation-types", timeout=10)
        
        if response.status_code == 200:
            data = response.json()
            print_success("AI generation types endpoint working!")
            print_info(f"Response: {json.dumps(data, indent=2)}")
            
            # Verify required fields
            required_fields = ['material_types', 'generation_types', 'difficulty_levels']
            for field in required_fields:
                if field not in data:
                    print_error(f"Missing required field: {field}")
                    return False
            
            # Verify material types
            expected_material_types = ['guide', 'checklist', 'template', 'whitepaper', 'report']
            for material_type in expected_material_types:
                if material_type not in data['material_types']:
                    print_error(f"Missing material type: {material_type}")
                    return False
            
            # Verify generation types
            expected_generation_types = ['summary', 'expanded_guide', 'checklist', 'case_study', 'best_practices']
            for gen_type in expected_generation_types:
                if gen_type not in data['generation_types']:
                    print_error(f"Missing generation type: {gen_type}")
                    return False
            
            # Verify difficulty levels
            expected_difficulties = ['beginner', 'intermediate', 'advanced']
            for difficulty in expected_difficulties:
                if difficulty not in data['difficulty_levels']:
                    print_error(f"Missing difficulty level: {difficulty}")
                    return False
            
            return True
        else:
            print_error(f"AI generation types failed with status {response.status_code}: {response.text}")
            return False
            
    except requests.exceptions.RequestException as e:
        print_error(f"AI generation types request failed: {str(e)}")
        return False

def test_ai_content_generation():
    """Test POST /api/ai/generate-content endpoint"""
    print_test_header("AI Content Generation")
    
    try:
        # Test data for content generation
        form_data = {
            'prompt': 'Create a comprehensive guide on network security fundamentals for beginners',
            'material_type': 'guide',
            'context': 'This is for cybersecurity training program',
            'title': 'Network Security Fundamentals Guide',
            'description': 'AI-generated guide covering basic network security concepts',
            'category': 'Security Guide',
            'tags': 'network, security, fundamentals, beginner',
            'save_as_resource': 'true'
        }
        
        response = requests.post(
            f"{API_BASE}/ai/generate-content",
            data=form_data,
            timeout=120  # AI generation may take longer
        )
        
        if response.status_code == 200:
            data = response.json()
            print_success("AI content generation successful!")
            print_info(f"Generated content length: {len(data.get('content', ''))}")
            print_info(f"Word count: {data.get('word_count', 'N/A')}")
            print_info(f"Resource ID: {data.get('resource_id', 'N/A')}")
            
            # Verify required fields
            required_fields = ['success', 'content', 'material_type', 'word_count', 'session_id']
            for field in required_fields:
                if field not in data:
                    print_error(f"Missing required field: {field}")
                    return False, None
            
            # Verify content is not empty
            if not data.get('content') or len(data['content']) < 100:
                print_error("Generated content is too short or empty")
                return False, None
            
            # Verify resource was saved
            if data.get('save_as_resource') and not data.get('resource_id'):
                print_error("Resource should be saved but no resource_id returned")
                return False, None
            
            return True, data.get('resource_id')
        else:
            print_error(f"AI content generation failed with status {response.status_code}: {response.text}")
            return False, None
            
    except requests.exceptions.RequestException as e:
        print_error(f"AI content generation request failed: {str(e)}")
        return False, None

def test_ai_generate_from_resource(resource_id):
    """Test POST /api/ai/generate-from-resource/{resource_id} endpoint"""
    print_test_header(f"AI Generate from Resource (ID: {resource_id})")
    
    try:
        form_data = {
            'generation_type': 'summary',
            'specific_prompt': 'Create a concise executive summary of this content',
            'save_as_new_resource': 'true',
            'new_title': 'Executive Summary - Network Security',
            'new_category': 'Summary'
        }
        
        response = requests.post(
            f"{API_BASE}/ai/generate-from-resource/{resource_id}",
            data=form_data,
            timeout=120
        )
        
        if response.status_code == 200:
            data = response.json()
            print_success("AI generation from resource successful!")
            print_info(f"Generated content length: {len(data.get('content', ''))}")
            print_info(f"Generation type: {data.get('generation_type')}")
            print_info(f"New resource ID: {data.get('new_resource_id', 'N/A')}")
            
            # Verify required fields
            required_fields = ['success', 'content', 'generation_type', 'source_resource_id', 'session_id']
            for field in required_fields:
                if field not in data:
                    print_error(f"Missing required field: {field}")
                    return False, None
            
            # Verify content is generated
            if not data.get('content') or len(data['content']) < 50:
                print_error("Generated content is too short or empty")
                return False, None
            
            return True, data.get('new_resource_id')
        elif response.status_code == 404:
            print_error("Resource not found for AI generation")
            return False, None
        else:
            print_error(f"AI generation from resource failed with status {response.status_code}: {response.text}")
            return False, None
            
    except requests.exceptions.RequestException as e:
        print_error(f"AI generation from resource request failed: {str(e)}")
        return False, None

def test_ai_create_quiz(resource_id):
    """Test POST /api/ai/create-quiz/{resource_id} endpoint"""
    print_test_header(f"AI Quiz Creation (ID: {resource_id})")
    
    try:
        form_data = {
            'question_count': '5',
            'difficulty': 'intermediate',
            'save_as_resource': 'true'
        }
        
        response = requests.post(
            f"{API_BASE}/ai/create-quiz/{resource_id}",
            data=form_data,
            timeout=120
        )
        
        if response.status_code == 200:
            data = response.json()
            print_success("AI quiz creation successful!")
            print_info(f"Quiz content length: {len(data.get('quiz_content', ''))}")
            print_info(f"Question count: {data.get('question_count')}")
            print_info(f"Difficulty: {data.get('difficulty')}")
            print_info(f"Quiz resource ID: {data.get('quiz_resource_id', 'N/A')}")
            
            # Verify required fields
            required_fields = ['success', 'quiz_content', 'question_count', 'difficulty', 'source_resource_id', 'session_id']
            for field in required_fields:
                if field not in data:
                    print_error(f"Missing required field: {field}")
                    return False, None
            
            # Verify quiz content
            if not data.get('quiz_content') or len(data['quiz_content']) < 100:
                print_error("Quiz content is too short or empty")
                return False, None
            
            # Verify question count matches
            if data.get('question_count') != 5:
                print_error(f"Expected 5 questions, got {data.get('question_count')}")
                return False, None
            
            return True, data.get('quiz_resource_id')
        elif response.status_code == 404:
            print_error("Resource not found for quiz creation")
            return False, None
        else:
            print_error(f"AI quiz creation failed with status {response.status_code}: {response.text}")
            return False, None
            
    except requests.exceptions.RequestException as e:
        print_error(f"AI quiz creation request failed: {str(e)}")
        return False, None

def test_ai_error_handling():
    """Test AI endpoints error handling"""
    print_test_header("AI Error Handling Tests")
    
    all_passed = True
    
    # Test 1: Generate from non-existent resource
    print_info("Testing generation from non-existent resource...")
    try:
        fake_id = str(uuid.uuid4())
        form_data = {
            'generation_type': 'summary',
            'specific_prompt': 'Create a summary'
        }
        
        response = requests.post(
            f"{API_BASE}/ai/generate-from-resource/{fake_id}",
            data=form_data,
            timeout=30
        )
        
        if response.status_code == 404:
            print_success("✅ Non-existent resource correctly returns 404")
        else:
            print_error(f"❌ Expected 404 for non-existent resource, got {response.status_code}")
            all_passed = False
            
    except requests.exceptions.RequestException as e:
        print_error(f"❌ Request failed: {str(e)}")
        all_passed = False
    
    # Test 2: Quiz from non-existent resource
    print_info("Testing quiz creation from non-existent resource...")
    try:
        fake_id = str(uuid.uuid4())
        form_data = {
            'question_count': '5',
            'difficulty': 'intermediate'
        }
        
        response = requests.post(
            f"{API_BASE}/ai/create-quiz/{fake_id}",
            data=form_data,
            timeout=30
        )
        
        if response.status_code == 404:
            print_success("✅ Non-existent resource for quiz correctly returns 404")
        else:
            print_error(f"❌ Expected 404 for non-existent resource quiz, got {response.status_code}")
            all_passed = False
            
    except requests.exceptions.RequestException as e:
        print_error(f"❌ Quiz request failed: {str(e)}")
        all_passed = False
    
    # Test 3: Content generation with missing required fields
    print_info("Testing content generation with missing fields...")
    try:
        form_data = {
            'prompt': 'Test prompt'
            # Missing title, description, category
        }
        
        response = requests.post(
            f"{API_BASE}/ai/generate-content",
            data=form_data,
            timeout=30
        )
        
        if response.status_code == 422:  # FastAPI validation error
            print_success("✅ Missing required fields correctly returns 422")
        else:
            print_error(f"❌ Expected 422 for missing fields, got {response.status_code}")
            all_passed = False
            
    except requests.exceptions.RequestException as e:
        print_error(f"❌ Content generation request failed: {str(e)}")
        all_passed = False
    
    return all_passed

def test_ai_integration_with_existing_resources():
    """Test AI integration with actual uploaded resources"""
    print_test_header("AI Integration with Existing Resources")
    
    # First upload a resource with content
    test_file_path = create_test_pdf()
    
    try:
        # Upload resource
        with open(test_file_path, 'rb') as f:
            files = {
                'file': ('cybersecurity_guide.pdf', f, 'application/pdf')
            }
            data = {
                'title': 'Cybersecurity Best Practices Guide',
                'description': 'Comprehensive guide on cybersecurity best practices for organizations',
                'category': 'Security Guide',
                'tags': 'cybersecurity, best practices, security, guide'
            }
            
            response = requests.post(
                f"{API_BASE}/admin/resources/upload",
                files=files,
                data=data,
                timeout=30
            )
        
        if response.status_code != 200:
            print_error(f"Failed to upload test resource: {response.status_code}")
            return False, []
        
        upload_data = response.json()
        resource_id = upload_data['id']
        print_success(f"Test resource uploaded with ID: {resource_id}")
        
        # Test AI generation from this resource
        ai_resource_ids = []
        
        # Test 1: Generate summary
        print_info("Testing summary generation...")
        summary_success, summary_id = test_ai_generate_from_resource(resource_id)
        if summary_success and summary_id:
            ai_resource_ids.append(summary_id)
        
        # Test 2: Create quiz
        print_info("Testing quiz creation...")
        quiz_success, quiz_id = test_ai_create_quiz(resource_id)
        if quiz_success and quiz_id:
            ai_resource_ids.append(quiz_id)
        
        # Verify AI-generated resources are in database
        for ai_id in ai_resource_ids:
            try:
                response = requests.get(f"{API_BASE}/resources/{ai_id}", timeout=10)
                if response.status_code == 200:
                    ai_resource = response.json()
                    print_success(f"✅ AI resource {ai_id} found in database")
                    print_info(f"Title: {ai_resource.get('title')}")
                    print_info(f"File type: {ai_resource.get('file_type')}")
                    print_info(f"Tags: {ai_resource.get('tags')}")
                else:
                    print_error(f"❌ AI resource {ai_id} not found in database")
            except Exception as e:
                print_error(f"❌ Error checking AI resource {ai_id}: {str(e)}")
        
        return True, [resource_id] + ai_resource_ids
        
    except Exception as e:
        print_error(f"AI integration test failed: {str(e)}")
        return False, []
    finally:
        try:
            os.unlink(test_file_path)
        except:
            pass

def run_all_tests():
    """Run all backend tests including Sprint 2 AI generation features"""
    print(f"{Colors.BOLD}{Colors.BLUE}")
    print("=" * 80)
    print("TRIAD AI-based CMS Backend Test Suite")
    print("Testing Sprint 2 AI Generation Features")
    print("=" * 80)
    print(f"{Colors.ENDC}")
    
    print_info(f"Testing backend at: {BASE_URL}")
    
    test_results = {}
    uploaded_resource_ids = []
    file_urls = []
    
    # 1. Health Check
    test_results['health_check'] = test_health_check()
    
    # 2. Get resources (empty)
    test_results['get_resources_empty'] = test_get_resources_empty()
    
    # 3. Get categories (empty)
    test_results['get_categories_empty'] = test_get_categories_empty()
    
    # 4. Test PDF upload (should work normally)
    print_info("Testing normal PDF upload...")
    upload_success, resource_id = test_file_upload()
    test_results['pdf_upload'] = upload_success
    if upload_success and resource_id:
        uploaded_resource_ids.append(resource_id)
    
    # 5. Test PPTX to PDF conversion
    print_info("Testing PPTX to PDF conversion...")
    pptx_success, pptx_resource_id = test_pptx_to_pdf_conversion()
    test_results['pptx_conversion'] = pptx_success
    if pptx_success and pptx_resource_id:
        uploaded_resource_ids.append(pptx_resource_id)
    
    # 6. Test file size limits
    test_results['file_size_limit'] = test_file_size_limit()
    
    # 7. Test invalid file types
    test_results['invalid_file_types'] = test_invalid_file_types()
    
    # 8. Test missing required fields
    test_results['missing_required_fields'] = test_missing_required_fields()
    
    # 9. Database verification for PPTX conversion
    if pptx_resource_id:
        test_results['database_verification'] = test_database_verification(pptx_resource_id)
    
    # 10. Get single resource tests
    for i, resource_id in enumerate(uploaded_resource_ids):
        test_name = f'get_single_resource_{i+1}'
        test_results[test_name] = test_get_single_resource(resource_id)
    
    # 11. Get resources after uploads
    test_results['get_resources_after_upload'] = test_get_resources_after_upload()
    
    # 12. Test file URL accessibility and content-type
    for i, resource_id in enumerate(uploaded_resource_ids):
        try:
            response = requests.get(f"{API_BASE}/resources/{resource_id}")
            if response.status_code == 200:
                resource_data = response.json()
                file_url = resource_data.get('file_url')
                if file_url:
                    file_urls.append(file_url)
                    test_name = f'file_url_accessibility_{i+1}'
                    test_results[test_name] = test_file_url_accessibility(file_url)
                    
                    test_name = f'pdf_content_type_{i+1}'
                    test_results[test_name] = test_pdf_content_type(file_url)
        except:
            pass
    
    # === SPRINT 2 AI GENERATION TESTS ===
    print(f"\n{Colors.BOLD}{Colors.YELLOW}")
    print("=" * 60)
    print("SPRINT 2 AI GENERATION TESTS")
    print("=" * 60)
    print(f"{Colors.ENDC}")
    
    # 13. AI Generation Types endpoint
    test_results['ai_generation_types'] = test_ai_generation_types()
    
    # 14. AI Content Generation
    ai_content_success, ai_resource_id = test_ai_content_generation()
    test_results['ai_content_generation'] = ai_content_success
    if ai_content_success and ai_resource_id:
        uploaded_resource_ids.append(ai_resource_id)
    
    # 15. AI Integration with existing resources
    integration_success, integration_resource_ids = test_ai_integration_with_existing_resources()
    test_results['ai_integration'] = integration_success
    if integration_success:
        uploaded_resource_ids.extend(integration_resource_ids)
    
    # 16. AI Error handling
    test_results['ai_error_handling'] = test_ai_error_handling()
    
    # 17. Delete resources (cleanup)
    for i, resource_id in enumerate(uploaded_resource_ids):
        test_name = f'delete_resource_{i+1}'
        test_results[test_name] = test_delete_resource(resource_id)
    
    # Print summary
    print(f"\n{Colors.BOLD}{Colors.BLUE}")
    print("=" * 80)
    print("TEST SUMMARY")
    print("=" * 80)
    print(f"{Colors.ENDC}")
    
    passed = 0
    total = 0
    failed_tests = []
    
    for test_name, result in test_results.items():
        total += 1
        if result:
            passed += 1
            print_success(f"{test_name}: PASSED")
        else:
            failed_tests.append(test_name)
            print_error(f"{test_name}: FAILED")
    
    print(f"\n{Colors.BOLD}Results: {passed}/{total} tests passed{Colors.ENDC}")
    
    if passed == total:
        print_success("🎉 All tests passed! Sprint 2 AI generation features working correctly!")
        return True
    else:
        print_error(f"❌ {total - passed} tests failed:")
        for test in failed_tests:
            print_error(f"  - {test}")
        return False

if __name__ == "__main__":
    success = run_all_tests()
    exit(0 if success else 1)